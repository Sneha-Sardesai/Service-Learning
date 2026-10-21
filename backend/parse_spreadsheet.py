"""Read uploaded spreadsheet bytes without executing formulas or macros."""

import csv
import datetime
import json
import sys
import zipfile
from io import BytesIO, StringIO


def json_value(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return value


def parse_csv(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for index, row in enumerate(csv.reader(StringIO(text), dialect)):
        if index >= 10002:
            raise ValueError("A worksheet may contain no more than 10,000 data rows.")
        if len(row) > 100:
            raise ValueError("A worksheet may contain no more than 100 columns.")
        rows.append(row)
    return [{"name": "CSV", "rows": rows}]


def parse_excel(content, extension):
    sheets = []
    if extension == ".xlsx":
        with zipfile.ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(item.file_size for item in members) > 100 * 1024 * 1024:
                raise ValueError("Workbook expands beyond the allowed 100 MB processing limit.")
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if len(workbook.worksheets) > 30:
            raise ValueError("A workbook may contain no more than 30 worksheets.")
        for sheet in workbook.worksheets:
            rows = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= 10002:
                    raise ValueError("A worksheet may contain no more than 10,000 data rows.")
                if len(row) > 100:
                    raise ValueError("A worksheet may contain no more than 100 columns.")
                rows.append([json_value(cell) for cell in row])
            sheets.append({"name": sheet.title, "rows": rows})
        workbook.close()
    else:
        import xlrd

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        if len(workbook.sheet_names()) > 30:
            raise ValueError("A workbook may contain no more than 30 worksheets.")
        for sheet in workbook.sheets():
            if sheet.nrows > 10002 or sheet.ncols > 100:
                raise ValueError("A worksheet may contain no more than 10,000 data rows and 100 columns.")
            rows = []
            for row_index in range(sheet.nrows):
                values = []
                for col in range(sheet.ncols):
                    cell = sheet.cell(row_index, col)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode).date().isoformat()
                    values.append(value)
                rows.append(values)
            sheets.append({"name": sheet.name, "rows": rows})
        workbook.release_resources()
    return sheets


def main():
    extension = sys.argv[1].lower()
    content = sys.stdin.buffer.read()
    sheets = parse_csv(content) if extension == ".csv" else parse_excel(content, extension)
    result = []
    for sheet in sheets:
        rows = sheet["rows"]
        while rows and not any(value not in (None, "") for value in rows[0]):
            rows.pop(0)
        if not rows:
            result.append({"name": sheet["name"], "headers": [], "rows": []})
            continue
        base_headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        headers = []
        used_headers = set()
        for index, header in enumerate(base_headers):
            candidate = header or f"Column {index + 1}"
            suffix = 2
            while candidate in used_headers:
                candidate = f"{header or f'Column {index + 1}'} ({suffix})"
                suffix += 1
            used_headers.add(candidate)
            headers.append(candidate)
        records = []
        for row in rows[1:10001]:
            if not any(value not in (None, "") for value in row):
                continue
            records.append({
                headers[index] or f"Column {index + 1}": json_value(row[index]) if index < len(row) else None
                for index in range(len(headers))
            })
        result.append({"name": sheet["name"], "headers": headers, "rows": records})
    print(json.dumps({"sheets": result}, ensure_ascii=True, default=str))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"error": f"{type(error).__name__}: {error}"}))
        sys.exit(1)
